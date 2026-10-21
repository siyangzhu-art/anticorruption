from pathlib import Path
import pandas as pd,json,zipfile,hashlib,shutil
from openpyxl import load_workbook
ROOT=Path(r'C:\Users\siyang zhu\Desktop\hometown\project_organized');O=ROOT/'result/promotion_connections_rolling_2026_10_20'
def digest(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  while b:=f.read(1024*1024):h.update(b)
 return h.hexdigest()
sources=[(Path(r'C:\Users\siyang zhu\Downloads\rerequestforaccesstothecped.zip'),'User-supplied CPED ZIP','user provided'),(ROOT/'data/raw/policy_implementation_search_20261019/JPE_policy_experimentation/data.zip','Wang and Yang, doi:10.7910/DVN/38SN2Y','CC0 1.0'),(ROOT/'data/raw/policy_implementation_search_20261019/JPE_policy_experimentation/code.zip','Wang and Yang, author mapping of promotion codes','CC0 1.0'),(ROOT/'result/city_outcomes_province_Local_2026_10_19/analysis_panel.csv.gz','Existing China City Statistical Yearbook panel','existing project source'),(ROOT/'result/inspection_hometown_design_2026_10_16/first_provincial_inspection_years_with_official_province_codes.csv','Verified 31-province first inspection calendar; source URLs in CSV','official public releases')]
for sub in ['cped_city_province_connections_20261019/leadership_roles_all.csv','cped_city_province_connections_20261019/workplace_classification.csv.gz','cped_city_province_connections_20261019/education_classification.csv','cped_city_province_connections_20261019/pair_first_connection_dates.csv','cped_pre_anticorruption_person_network_20261019/person_edges_pre2012.csv.gz','cped_pre_anticorruption_person_network_20261019/school_eligibility_pre2012.csv.gz','cped_pre_anticorruption_person_network_20261019/date_ambiguous_pairs_pre2012.csv.gz']:sources.append((ROOT/'result'/sub,'Derived solely from user CPED package','user provided; original preserved'))
pd.DataFrame([dict(path=str(p),source=s,license_or_access=l,bytes=p.stat().st_size,sha256=digest(p)) for p,s,l in sources]).to_csv(O/'input_manifest.csv',index=False,encoding='utf-8-sig')
v=json.loads((O/'validation_summary.json').read_text(encoding='utf8'));e=json.loads((O/'explorer_validation.json').read_text(encoding='utf8'));r=pd.read_csv(O/'all_estimates_validated.csv');statuses=pd.read_csv(O/'model_status.csv')
w=load_workbook(O/'省市关系与晋升_全部结果.xlsx',read_only=True)
assert w['全部系数'].max_row-1==len(r);assert w['全部模型及留空'].max_row-1==len(statuses)
assert len(list((O/'figures').glob('*.png')))==28
assert not list(O.glob('*.pdf'));assert v['all_passed'] and e['status']=='PASS'
q={'status':'PASS','workbook_sheets':len(w.sheetnames),'coefficient_rows':len(r),'models':len(statuses),'figures':28,'numerical_validation':v,'explorer_validation':e,'no_pdf':True,'window_coverage_is_explicit':bool(r.loc[r.kind.eq('rolling'),'observed_time_points'].notna().all())};w.close()
(O/'delivery_validation.json').write_text(json.dumps(q,ensure_ascii=False,indent=2),encoding='utf8');shutil.copy2(__file__,O/'code/delivery_qc.py')
files=[dict(path=str(p.relative_to(O)),bytes=p.stat().st_size,sha256=digest(p)) for p in O.rglob('*') if p.is_file() and p.name!='file_manifest.csv'];pd.DataFrame(files).to_csv(O/'file_manifest.csv',index=False,encoding='utf-8-sig')
archive=O.parent/'省市官员关系与晋升_滚动回归_数据图表代码.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
 for p in O.rglob('*'):
  if p.is_file():z.write(p,str(p.relative_to(O)))
with zipfile.ZipFile(archive) as z:assert z.testzip() is None;count=len(z.namelist())
print(json.dumps({**q,'archive_files':count,'archive_bytes':archive.stat().st_size,'archive_sha256':digest(archive)},ensure_ascii=False),flush=True)
